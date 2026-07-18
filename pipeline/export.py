"""Project scoring results into output tables and export them.

Two layers:

* **Legacy** (unchanged): ``COLUMNS`` + ``build_dataframe`` + ``to_csv_bytes`` + ``to_xlsx_bytes``.
* **Release 0.3.1 lead-generator workbook**: a three-sheet XLSX modelled on the team's reference
  workbook (``data/reference/fintech_scored_output_example.xlsx``) —
  ``Fintech Leads Scored`` (the daily working sheet) · ``AI Details`` (audit/technical) · ``Summary``.

The working sheet stays compact (single-height rows, no wrapped technical columns); all AI audit
fields live on a secondary sheet. Reviewer fields are always emitted empty; no lead is auto-approved;
no Google Sheets API, no chain-of-thought, no raw JSON is exported.
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook  # noqa: F401  (load_workbook used by tests)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from scoring import Lead, ScoringResult

# ===========================================================================
# Legacy flat table (unchanged — backward compatible)
# ===========================================================================

COLUMNS = [
    "#", "Lead Score", "Priority Status",
    "First Name", "Last Name", "Job Title", "Job Started",
    "LinkedIn URL", "Number of Connections", "Location",
    "Company", "Company LinkedIn URL", "Company Website",
    "LinkedIn Employees", "LinkedIn Founded Year", "LinkedIn Industry",
    "LinkedIn Specialities",
    "ICP Signals", "Score Reason", "Score Breakdown",
    "Confidence", "Unknowns",
]


def _score_breakdown(result: ScoringResult) -> str:
    labels = {"title": "Title", "industry": "Industry", "company_size": "Size",
              "location": "Loc", "signals": "Signals"}
    return " ".join(f"{labels.get(k, k)}:{d.points}" for k, d in result.dimensions.items())


def build_dataframe(pairs: list[tuple[Lead, ScoringResult]]) -> pd.DataFrame:
    """LEGACY canonical DataFrame. Preserved for backward compatibility."""
    rows = []
    for i, (lead, result) in enumerate(pairs, start=1):
        f = lead.fields
        rows.append({
            "#": i, "Lead Score": result.score, "Priority Status": result.category,
            "First Name": f.get("first_name", ""), "Last Name": f.get("last_name", ""),
            "Job Title": f.get("job_title", ""), "Job Started": f.get("job_started", ""),
            "LinkedIn URL": f.get("linkedin_url", ""), "Number of Connections": f.get("connections", ""),
            "Location": f.get("location", ""), "Company": f.get("company", ""),
            "Company LinkedIn URL": f.get("company_linkedin_url", ""),
            "Company Website": f.get("company_website", ""),
            "LinkedIn Employees": f.get("company_size_range", ""),
            "LinkedIn Founded Year": f.get("founded_year", ""),
            "LinkedIn Industry": f.get("industry", ""),
            "LinkedIn Specialities": f.get("specialities", ""),
            "ICP Signals": "; ".join(result.signals), "Score Reason": result.reason,
            "Score Breakdown": _score_breakdown(result), "Confidence": result.confidence,
            "Unknowns": "; ".join(result.unknowns),
        })
    return pd.DataFrame(rows, columns=COLUMNS)


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


def to_xlsx_bytes(df: pd.DataFrame, sheet_name: str = "Leads") -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name[:31] or "Leads")
    return buf.getvalue()


# ===========================================================================
# Release 0.3.1 — lead-generator workbook (3 sheets)
# ===========================================================================

SHEET_MAIN = "Fintech Leads Scored"
SHEET_AI = "AI Details"
SHEET_SUMMARY = "Summary"
REVIEW_STATUS_DEFAULT = "Pending"

# Sheet 1 — the daily working sheet. User-facing "Priority" (operational label) replaces the
# internal category on this sheet only.
MAIN_COLUMNS = [
    "#", "Lead Score", "Priority", "First Name", "Last Name", "Job Title", "Job Started",
    "LinkedIn URL", "Number of Connections", "Location", "Company", "Company LinkedIn URL",
    "Company Website", "LinkedIn Employees", "LinkedIn Founded Year", "LinkedIn Industry",
    "LinkedIn Specialities", "ICP Signals", "Score Reason", "Score Breakdown",
    "Review Status", "Human Decision", "Rejection Reason", "Reviewer Comment",
]

# Sheet 2 — technical / audit fields; preserves Raw ICP Score, Operational Lead Score and the
# internal Decision-Layer category alongside coverage/confidence/unknowns.
AI_COLUMNS = [
    "#", "LinkedIn URL", "Internal Decision", "Raw ICP Score", "Operational Lead Score",
    "Evidence-Adjusted Fit", "Data Coverage", "Decision Confidence", "Evidence Summary",
    "Confirmed Dealbreakers", "Suspected Dealbreakers", "Unknown Fields", "Confidence Reasons",
    "Validation Warnings", "Review Recommendation", "Dealbreaker State", "Model Confidence",
    "Mock Result",
]

# Operational lead score + priority are computed by the Decision Layer (Python is the sole
# authority); the workbook displays them directly. Sort/summary use this fixed ordering.
_OP_ORDER = ["Priority 1", "Priority 2", "Priority 3", "Priority 4", "Priority 5", "Disqualified"]
_OP_RANK = {label: i for i, label in enumerate(_OP_ORDER)}


def _priority_of(result) -> str:
    return getattr(result, "operational_priority", None) or result.category


def _op_score_of(result) -> int:
    v = getattr(result, "operational_lead_score", None)
    if v is not None:
        return v
    return result.score if result.score is not None else 0

_MAIN_URLS = {"LinkedIn URL", "Company LinkedIn URL", "Company Website"}
_MAIN_WRAP = {"Score Reason", "Score Breakdown", "Reviewer Comment"}
_MAIN_CENTER = {"#", "Lead Score"}
# Compact but sufficient widths. _write_table also enforces width >= len(header)+2 so no header
# is ever truncated by a too-small configured width.
_MAIN_WIDTHS = {
    "#": 5, "Lead Score": 11, "Priority": 14, "First Name": 14, "Last Name": 14,
    "Job Title": 30, "Job Started": 12, "LinkedIn URL": 36, "Number of Connections": 22,
    "Location": 26, "Company": 22, "Company LinkedIn URL": 24, "Company Website": 26,
    "LinkedIn Employees": 19, "LinkedIn Founded Year": 22, "LinkedIn Industry": 22,
    "LinkedIn Specialities": 30, "ICP Signals": 26, "Score Reason": 46, "Score Breakdown": 28,
    "Review Status": 14, "Human Decision": 16, "Rejection Reason": 18, "Reviewer Comment": 22,
}

# Priority cell styling by OPERATIONAL label: (fill hex, font hex, bold). Label text always visible.
_PRIORITY_STYLE = {
    "Priority 1": ("1E7B34", "FFFFFF", True),     # dark green / white bold
    "Priority 2": ("C6EFCE", "1E4620", True),     # light green / dark bold
    "Priority 3": ("FFD966", "3F3000", True),     # gold / dark bold
    "Priority 4": ("BDD7EE", "1F3864", False),    # light blue / dark
    "Priority 5": ("D9D9D9", "3F3F3F", False),    # light gray / dark
    "Disqualified": ("C00000", "FFFFFF", True),   # red / white bold
}
_HEADER_FILL = "F2F2F2"

# Subtle neutral-gray borders; header row gets a stronger bottom edge.
_THIN = Side(style="thin", color="BFBFBF")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_BORDER = Border(left=_THIN, right=_THIN, top=_THIN,
                        bottom=Side(style="medium", color="808080"))


def sort_pairs(pairs: list[tuple[Lead, ScoringResult]]) -> list[tuple[Lead, ScoringResult]]:
    """Sort for the working sheet: operational priority rank, then Lead Score desc, then Company
    asc, then Last Name asc. Disqualified always sinks to the bottom. Deterministic — main and AI
    sheets share this order so ``#`` cross-refs."""
    def key(pair):
        lead, result = pair
        op = _priority_of(result)
        return (_OP_RANK.get(op, 4), -int(_op_score_of(result) or 0),
                (lead.fields.get("company") or "").lower(), (lead.fields.get("last_name") or "").lower())
    return sorted(pairs, key=key)
_AI_URLS = {"LinkedIn URL"}
_AI_WRAP = {"Evidence Summary", "Confirmed Dealbreakers", "Suspected Dealbreakers", "Unknown Fields",
            "Confidence Reasons", "Validation Warnings"}
_AI_WIDTHS = {
    "#": 5, "LinkedIn URL": 34, "Internal Decision": 22, "Raw ICP Score": 13,
    "Operational Lead Score": 20, "Evidence-Adjusted Fit": 16,
    "Data Coverage": 12, "Decision Confidence": 15, "Evidence Summary": 50, "Confirmed Dealbreakers": 34,
    "Suspected Dealbreakers": 34, "Unknown Fields": 30, "Confidence Reasons": 40,
    "Validation Warnings": 40, "Review Recommendation": 18, "Dealbreaker State": 15,
    "Model Confidence": 15, "Mock Result": 13,
}


def _pct(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value
    return f"{int(value)}%"


def _evidence_summary(result: ScoringResult) -> str:
    parts = [f"{name}: {d.evidence}" for name, d in result.dimensions.items() if d.evidence]
    return " | ".join(parts)[:500]


def _get(result, attr, default):
    v = getattr(result, attr, default)
    return default if v is None else v


def build_main_rows(pairs: list[tuple[Lead, ScoringResult]]) -> list[dict]:
    rows = []
    for i, (lead, result) in enumerate(sort_pairs(pairs), start=1):
        f = lead.fields
        op_score = _op_score_of(result)
        rows.append({
            "#": i,
            "Lead Score": int(op_score),                          # OPERATIONAL score (0 if excluded)
            "Priority": _priority_of(result),                     # OPERATIONAL priority (from score)
            "First Name": f.get("first_name", ""), "Last Name": f.get("last_name", ""),
            "Job Title": f.get("job_title", ""), "Job Started": f.get("job_started", ""),
            "LinkedIn URL": f.get("linkedin_url", ""), "Number of Connections": f.get("connections", ""),
            "Location": f.get("location", ""), "Company": f.get("company", ""),
            "Company LinkedIn URL": f.get("company_linkedin_url", ""),
            "Company Website": f.get("company_website", ""),
            "LinkedIn Employees": f.get("company_size_range", ""),
            "LinkedIn Founded Year": f.get("founded_year", ""),
            "LinkedIn Industry": f.get("industry", ""),
            "LinkedIn Specialities": f.get("specialities", ""),
            "ICP Signals": "; ".join(result.signals or []),
            "Score Reason": result.reason, "Score Breakdown": _score_breakdown(result),
            "Review Status": REVIEW_STATUS_DEFAULT,
            "Human Decision": "", "Rejection Reason": "", "Reviewer Comment": "",
        })
    return rows


def build_ai_rows(pairs: list[tuple[Lead, ScoringResult]]) -> list[dict]:
    rows = []
    for i, (lead, result) in enumerate(sort_pairs(pairs), start=1):
        rows.append({
            "#": i, "LinkedIn URL": lead.fields.get("linkedin_url", ""),
            "Internal Decision": getattr(result, "internal_category", None) or result.category,
            "Raw ICP Score": result.raw_icp_score if result.raw_icp_score is not None else result.score,
            "Operational Lead Score": _op_score_of(result),
            "Evidence-Adjusted Fit": _pct(result.evidence_adjusted_fit),
            "Data Coverage": _pct(result.evidence_coverage),
            "Decision Confidence": _get(result, "decision_confidence", ""),
            "Evidence Summary": _evidence_summary(result),
            "Confirmed Dealbreakers": "; ".join(_get(result, "confirmed_dealbreakers", [])),
            "Suspected Dealbreakers": "; ".join(_get(result, "suspected_dealbreakers", [])),
            "Unknown Fields": "; ".join(result.unknowns or []),
            "Confidence Reasons": "; ".join(_get(result, "confidence_reasons", [])),
            "Validation Warnings": "; ".join(_get(result, "validation_warnings", [])),
            "Review Recommendation": _get(result, "review_recommendation", ""),
            "Dealbreaker State": _get(result, "dealbreaker_state", ""),
            "Model Confidence": _get(result, "model_confidence", ""),
            "Mock Result": "MOCK/OFFLINE" if getattr(result, "is_mock", False) else "real",
        })
    return rows


def build_main_dataframe(pairs) -> pd.DataFrame:
    return pd.DataFrame(build_main_rows(pairs), columns=MAIN_COLUMNS)


def build_ai_dataframe(pairs) -> pd.DataFrame:
    return pd.DataFrame(build_ai_rows(pairs), columns=AI_COLUMNS)


def build_summary_rows(pairs, campaign: str, icp_name: str, generated_at: str):
    total = len(pairs)
    results = [r for _, r in pairs]
    success = sum(1 for r in results if r.error is None)
    internal = Counter(getattr(r, "internal_category", None) or r.category for r in results)
    op_counts = Counter(_priority_of(r) for r in results)
    excluded = op_counts.get("Disqualified", 0)
    qualified = success - excluded
    mock = sum(1 for r in results if getattr(r, "is_mock", False))

    rows = [
        ("Campaign", campaign), ("ICP", icp_name), ("Generated", generated_at),
        ("", ""), ("Total processed", total), ("Qualified", qualified),
        ("Disqualified / Excluded", excluded), ("Successful", success), ("Failed", total - success),
        ("", ""), ("Priority distribution", ""),
    ]
    for label in _OP_ORDER:
        rows.append((f"  {label}", op_counts.get(label, 0)))
    rows.append(("", ""))
    rows.append(("Internal category distribution", ""))
    for label, count in internal.most_common():
        rows.append((f"  {label}", count))
    rows += [
        ("", ""),
        ("Review — Pending", total), ("Review — Approved", 0),
        ("Review — Rejected", 0), ("Review — Skipped", 0),
        ("", ""), ("Real results", total - mock), ("Mock results", mock),
        ("", ""),
        ("Note", "Human approval is required before any outreach. Filter Human Decision = Approved "
                 "and copy LinkedIn URLs into Linked Helper manually."),
    ]
    return rows


# --- CSV fallbacks ----------------------------------------------------------

def to_main_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


def to_ai_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


# --- workbook ---------------------------------------------------------------

def _write_table(ws, columns, rows, widths, url_cols, wrap_cols, *, freeze, autofilter=False,
                 center_cols=frozenset(), style_priority=False):
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    link_font = Font(color="0563C1", underline="single")
    top = Alignment(vertical="top")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")

    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = _HEADER_BORDER                      # stronger bottom edge on the header row
        # never let a too-small configured width truncate the header text
        ws.column_dimensions[cell.column_letter].width = max(widths.get(name, 16), len(str(name)) + 2)

    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(columns, start=1):
            value = row.get(name, "")
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = _CELL_BORDER                    # full used-range grid
            if name in center_cols:
                cell.alignment = center
            elif name in wrap_cols:
                cell.alignment = wrap
            else:
                cell.alignment = top
            if name in url_cols and isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.font = link_font
            if style_priority and name == "Priority":
                style = _PRIORITY_STYLE.get(str(value))
                if style:
                    fill_hex, font_hex, bold = style
                    cell.fill = PatternFill("solid", fgColor=fill_hex)
                    cell.font = Font(color=font_hex, bold=bold)

    ws.freeze_panes = freeze
    if autofilter:
        last_col = ws.cell(row=1, column=len(columns)).column_letter
        ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"


def _write_summary(ws, summary_rows):
    for col, name in ((1, "Metric"), (2, "Value")):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.border = _HEADER_BORDER
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 66
    bold_metrics = ("Campaign", "ICP", "Priority distribution", "Internal category distribution",
                    "Total processed")
    for r, (metric, value) in enumerate(summary_rows, start=2):
        a = ws.cell(row=r, column=1, value=metric)
        b = ws.cell(row=r, column=2, value=value)
        a.font = Font(bold=metric in bold_metrics)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        # border only POPULATED cells (leave spacer rows unstyled)
        if metric not in ("", None):
            a.border = _CELL_BORDER
        if value not in ("", None):
            b.border = _CELL_BORDER


def to_workbook_bytes(pairs, icp_name: str, campaign: str | None = None,
                      generated_at: str | None = None) -> bytes:
    """Build the three-sheet lead-generator workbook: Fintech Leads Scored / AI Details / Summary."""
    campaign = campaign or icp_name
    generated_at = generated_at or datetime.now().strftime("%Y-%m-%d %H:%M")
    return workbook_bytes_from_rows(build_main_rows(pairs), build_ai_rows(pairs),
                                    build_summary_rows(pairs, campaign, icp_name, generated_at))


def workbook_bytes_from_rows(main_rows: list[dict], ai_rows: list[dict], summary_rows) -> bytes:
    """ADDITIVE (Sprint 13b): build the same three-sheet workbook from already-assembled row dicts.

    The canonical schema is unchanged — rows must be keyed by ``MAIN_COLUMNS`` / ``AI_COLUMNS``. This
    lets a caller that already holds a fully-assembled view model (e.g. the Human Review projection)
    serialize it directly, without rehydrating legacy Lead/ScoringResult objects just to satisfy the
    pairs-based signature. ``to_workbook_bytes`` delegates here, so both paths produce identical output."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = SHEET_MAIN
    # working sheet: sorted best-first, priority-coloured, no filter dropdowns, freeze D2,
    # centred numeric Lead Score, wrapping only in the explanation columns.
    _write_table(ws1, MAIN_COLUMNS, main_rows, _MAIN_WIDTHS, _MAIN_URLS, _MAIN_WRAP,
                 freeze="D2", autofilter=False, center_cols=_MAIN_CENTER, style_priority=True)
    ws2 = wb.create_sheet(SHEET_AI)
    _write_table(ws2, AI_COLUMNS, ai_rows, _AI_WIDTHS, _AI_URLS, _AI_WRAP, freeze="B2", autofilter=False)
    ws3 = wb.create_sheet(SHEET_SUMMARY)
    _write_summary(ws3, summary_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
