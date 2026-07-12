"""Offline tests for the lead-generator workbook export (pipeline/export.py, Release 0.3.4).

No network, no API, no pytest dependency:
    ./.venv/bin/python tests/test_export.py
"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "pipeline"))

import pandas as pd                    # noqa: E402
from openpyxl import load_workbook     # noqa: E402
import scoring as sc                   # noqa: E402
import decision as dc                  # noqa: E402
import export                          # noqa: E402

_TECHNICAL = {"Data Coverage", "Evidence-Adjusted Fit", "Decision Confidence", "Confidence",
              "Unknown Fields", "Confidence Reasons", "Validation Warnings", "Dealbreaker State",
              "Model Confidence", "Mock Result", "Evidence Summary", "Internal Decision",
              "Raw ICP Score", "Operational Lead Score"}


def _lead(idx, **fields):
    return sc.Lead(index=idx, fields=fields, raw={}, previous_roles=[])


def _result(idx, *, internal, raw, db_state="none", coverage=60, fit=83, conf=55,
            confidence="medium", confirmed=None, suspected=None, unknowns=None, is_mock=False):
    op_score = 0 if db_state == "confirmed" else raw
    op_priority = "Disqualified" if db_state == "confirmed" else dc.operational_priority(raw)
    return sc.ScoringResult(
        lead_index=idx, icp="FinTech", score=op_score, category=op_priority,
        dimensions={"Subsegment fit": sc.Dimension(20, 25, "wealth-tech platform")},
        hard_dealbreaker=(db_state == "confirmed"), dealbreaker_reason=None,
        reason="Sample qualification reason.", signals=["1000 connections"], confidence=confidence,
        unknowns=unknowns or [], model="claude-haiku",
        raw_icp_score=raw, operational_lead_score=op_score, operational_priority=op_priority,
        internal_category=internal, provisional_priority=op_priority, evidence_coverage=coverage,
        evidence_adjusted_fit=fit, decision_confidence=conf, dealbreaker_state=db_state,
        confirmed_dealbreakers=confirmed or [], suspected_dealbreakers=suspected or [],
        review_recommendation="standard_review", confidence_reasons=[f"coverage {coverage}%"],
        validation_warnings=["enrichment guard: 'Reachability' removed"], model_confidence="low",
        is_mock=is_mock)


def _pairs():
    return [
        (_lead(0, first_name="Aaron", last_name="Wormus", job_title="CTO", company="SMArtX",
               linkedin_url="https://www.linkedin.com/in/aaron", location="FL, US",
               company_size_range="51-200", industry="Financial Services", connections="2206",
               founded_year="2018", specialities="WealthTech",
               company_linkedin_url="https://www.linkedin.com/company/smartx",
               company_website="https://smartx.com", job_started="03/2017"),
         _result(0, internal="A / High", raw=78, db_state="suspected", suspected=["no hiring signal"],
                 unknowns=["Stage & funding fit"])),
        (_lead(1, first_name="Ditrick", last_name="Dunn", job_title="BD", company="Servicing Co",
               linkedin_url="https://www.linkedin.com/in/ditrickdunn", location="US"),
         _result(1, internal="Disqualified", raw=32, db_state="confirmed",
                 confirmed=["Dev shop / BPO"], unknowns=["Reachability"])),
        (_lead(2, first_name="Mock", last_name="Lead", job_title="CEO", company="Placeholder",
               linkedin_url="https://www.linkedin.com/in/mock", location="US"),
         _result(2, internal="B / Normal", raw=45, is_mock=True)),
    ]


def _mapping_pairs():
    cats = ["A+ / Hot", "A+ Candidate — Enrichment Required", "A / High", "B / Normal",
            "C / Low", "Not Relevant", "Disqualified"]
    return [(_lead(i, first_name="F", last_name=f"L{i}", company=f"C{i}",
                   linkedin_url=f"https://www.linkedin.com/in/{i}"),
             _result(i, internal=c, raw=50)) for i, c in enumerate(cats)]


def _sort_fixture():
    def mk(idx, score, company, last, db="none"):
        return (_lead(idx, first_name="F", last_name=last, company=company,
                      linkedin_url=f"https://www.linkedin.com/in/{idx}"),
                _result(idx, internal="A+ / Hot", raw=score, db_state=db,
                        confirmed=["x"] if db == "confirmed" else None))
    return [
        mk(0, 92, "Zeta", "Zed"),                    # P1
        mk(1, 80, "Mid", "Mm"),                      # P2
        mk(2, 92, "Alpha", "Bravo"),                 # P1
        mk(3, 90, "Xco", "Xx", db="confirmed"),      # high raw but Disqualified -> bottom
        mk(4, 91, "Alpha", "Alfa"),                  # P1
        mk(5, 46, "Beta", "Cand"),                   # P4
    ]


def _wb(pairs=None):
    return load_workbook(io.BytesIO(export.to_workbook_bytes(pairs or _pairs(), "FinTech")))


def _sv(rows, key):
    for k, v in rows:
        if k == key:
            return v
    return None


def _col(ws, name):
    return [c.value for c in ws[1]].index(name) + 1


# --- structure --------------------------------------------------------------

def test_exactly_three_sheet_names():
    assert _wb().sheetnames == ["Fintech Leads Scored", "AI Details", "Summary"]


def test_main_column_order_uses_priority():
    header = [c.value for c in _wb()["Fintech Leads Scored"][1]]
    assert header == export.MAIN_COLUMNS and "Priority" in header and len(header) == 24


def test_main_displays_operational_priority_and_score():
    mdf = export.build_main_dataframe(_pairs())
    # top row is the highest operational lead (A/High raw 78 -> Priority 2)
    assert mdf.iloc[0]["Priority"] == "Priority 2" and mdf.iloc[0]["Lead Score"] == 78


def test_technical_fields_absent_from_main_sheet():
    assert not (set(c.value for c in _wb()["Fintech Leads Scored"][1]) & _TECHNICAL)


def test_ai_details_columns_and_internal_decision():
    adf = export.build_ai_dataframe(_mapping_pairs())
    assert list(adf.columns) == export.AI_COLUMNS
    for col in ("Internal Decision", "Raw ICP Score", "Operational Lead Score"):
        assert col in adf.columns
    assert "A+ Candidate — Enrichment Required" in set(adf["Internal Decision"])


def test_main_sheet_never_shows_enrichment_required():
    ws = _wb(_mapping_pairs())["Fintech Leads Scored"]
    for row in ws.iter_rows(values_only=True):
        assert not any("Enrichment" in str(v) for v in row if v)
    assert not any("Enrichment" in str(v) for v in export.build_main_dataframe(_mapping_pairs())["Priority"])


# --- sorting ----------------------------------------------------------------

def test_priority_first_then_score_desc():
    mdf = export.build_main_dataframe(_sort_fixture())
    assert list(mdf["Priority"]) == ["Priority 1", "Priority 1", "Priority 1",
                                     "Priority 2", "Priority 4", "Disqualified"]
    assert list(mdf["Lead Score"][:3]) == [92, 92, 91]          # score desc within Priority 1
    assert mdf.iloc[2]["Lead Score"] == 91 and mdf.iloc[3]["Lead Score"] == 80  # low P1 above high P2


def test_disqualified_rows_last():
    mdf = export.build_main_dataframe(_sort_fixture())
    assert mdf.iloc[-1]["Priority"] == "Disqualified" and mdf.iloc[-1]["Lead Score"] == 0


def test_company_then_lastname_tertiary_sort():
    mdf = export.build_main_dataframe(_sort_fixture())
    assert list(mdf["Last Name"][:2]) == ["Bravo", "Zed"]       # both 92, company Alpha before Zeta


# --- confirmed-exclusion score contract -------------------------------------

def test_confirmed_exclusion_shows_zero_on_main_and_raw_on_ai():
    pairs = _pairs()
    mdf = export.build_main_dataframe(pairs)
    disq = mdf[mdf["Priority"] == "Disqualified"].iloc[0]
    assert disq["Lead Score"] == 0                              # main shows operational 0
    adf = export.build_ai_dataframe(pairs)
    ai_disq = adf[adf["Internal Decision"] == "Disqualified"].iloc[0]
    assert ai_disq["Raw ICP Score"] == 32 and ai_disq["Operational Lead Score"] == 0


# --- styling & borders ------------------------------------------------------

def test_priority_cells_styled():
    ws = _wb()["Fintech Leads Scored"]                          # sorted: P2(78), P4(45), Disq(0)
    pc = _col(ws, "Priority")
    p2, p4, disq = ws.cell(row=2, column=pc), ws.cell(row=3, column=pc), ws.cell(row=4, column=pc)
    assert p2.value == "Priority 2" and p2.fill.fgColor.rgb.endswith("C6EFCE") and p2.font.bold
    assert p4.value == "Priority 4" and p4.fill.fgColor.rgb.endswith("BDD7EE")
    assert disq.value == "Disqualified" and disq.fill.fgColor.rgb.endswith("C00000")
    assert disq.font.color.rgb.endswith("FFFFFF") and disq.font.bold


def test_populated_cells_have_borders():
    c = _wb()["Fintech Leads Scored"].cell(row=2, column=1)
    assert all(getattr(c.border, s).style == "thin" for s in ("left", "right", "top", "bottom"))


def test_header_bottom_border_stronger():
    assert _wb()["Fintech Leads Scored"].cell(row=1, column=1).border.bottom.style == "medium"


def test_empty_rows_outside_used_range_not_styled():
    ws = _wb()["Fintech Leads Scored"]
    empty = ws.cell(row=ws.max_row + 2, column=1)
    assert empty.border.left.style is None and empty.border.bottom.style is None


def test_summary_spacer_cells_not_styled():
    ws = _wb()["Summary"]
    found = False
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value in ("", None) and ws.cell(row=r, column=2).value in ("", None):
            assert ws.cell(row=r, column=1).border.left.style is None
            found = True
    assert found


# --- preserved rules --------------------------------------------------------

def test_freeze_panes_d2_and_no_autofilter():
    ws = _wb()["Fintech Leads Scored"]
    assert ws.freeze_panes == "D2" and ws.auto_filter.ref is None


def test_headers_not_truncated_by_width():
    ws = _wb()["Fintech Leads Scored"]
    for c, name in enumerate(export.MAIN_COLUMNS, start=1):
        assert ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width >= len(name), name


def test_lead_score_numeric():
    ws = _wb()["Fintech Leads Scored"]
    sc_col = _col(ws, "Lead Score")
    for r in range(2, ws.max_row + 1):
        assert isinstance(ws.cell(row=r, column=sc_col).value, int)


def test_reviewer_fields_blank_and_status_pending():
    mdf = export.build_main_dataframe(_pairs())
    assert (mdf["Review Status"] == "Pending").all()
    for col in ("Human Decision", "Rejection Reason", "Reviewer Comment"):
        assert (mdf[col] == "").all()


def test_urls_intact_and_clickable():
    ws = _wb()["Fintech Leads Scored"]
    cell = ws.cell(row=2, column=_col(ws, "LinkedIn URL"))
    assert cell.value.startswith("https://www.linkedin.com/in/") and cell.hyperlink is not None


def test_summary_operational_and_internal_distribution():
    rows = export.build_summary_rows(_pairs(), "FinTech Campaign", "FinTech", "2026-07-12 00:00")
    assert _sv(rows, "Campaign") == "FinTech Campaign" and _sv(rows, "ICP") == "FinTech"
    assert _sv(rows, "  Priority 2") == 1 and _sv(rows, "  Priority 4") == 1
    assert _sv(rows, "  Disqualified") == 1
    assert _sv(rows, "  A / High") == 1                         # internal distribution retained
    assert _sv(rows, "Total processed") == 3 and _sv(rows, "Disqualified / Excluded") == 1


def test_legacy_csv_and_dataframe_compatible():
    df = export.build_dataframe(_pairs())
    assert list(df.columns) == export.COLUMNS and len(df) == 3
    reread = pd.read_csv(io.BytesIO(export.to_main_csv_bytes(export.build_main_dataframe(_pairs()))))
    assert list(reread.columns) == export.MAIN_COLUMNS and len(reread) == 3


# --- runner -----------------------------------------------------------------

def _run():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    passed = 0
    for t in tests:
        try:
            t()
            print(f"  PASS  {t.__name__}")
            passed += 1
        except AssertionError as e:
            print(f"  FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa: BLE001
            import traceback
            print(f"  ERROR {t.__name__}: {type(e).__name__}: {e}")
            traceback.print_exc()
    print(f"\n{passed}/{len(tests)} passed")
    return passed == len(tests)


if __name__ == "__main__":
    sys.exit(0 if _run() else 1)
