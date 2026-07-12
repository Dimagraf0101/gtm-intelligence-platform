"""Release 0.3 Internal Alpha — end-to-end acceptance test (20 FinTech leads, real API).

Single run. Exercises ICP extraction -> profile -> normalization -> evidence -> real-model
proposals -> Python decision -> enrichment guard -> confidence/coverage -> 3-sheet workbook +
CSV fallbacks. Writes 5 output files to a fresh timestamped folder. No new features; harness only.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "pipeline"))
csv.field_size_limit(10_000_000)

import config                      # noqa: E402  loads .env
import evidence as ev              # noqa: E402
import icp_profile as ip           # noqa: E402
import scoring as sc               # noqa: E402
import export                      # noqa: E402
from icp_pdf import extract_icp    # noqa: E402
from openpyxl import load_workbook  # noqa: E402

RAW = ROOT / "data" / "raw" / "fintech_raw_vayne.csv"
BENCH = ROOT / "data" / "benchmarks" / "legacy" / "fintech_manual_v2_unvalidated.csv"
ICP_PDF = ROOT / "icp" / "Innotechfy_Fintech_Playbook.pdf"
MODEL = "claude-haiku-4-5-20251001"
P_IN, P_OUT, P_CR, P_CW = 1.00, 5.00, 0.10, 1.25
PROJECTION_N = 20_000


def slug(u):
    u = (u or "").strip().lower().split("?")[0].rstrip("/")
    m = re.search(r"linkedin\.com/in/([^/]+)", u)
    return m.group(1) if m else None


class CapturingClient(sc.AnthropicClient):
    def __init__(self):
        super().__init__(model=MODEL)
        self.usages = []

    def complete(self, system, user):
        try:
            m = self._client.messages.create(model=self.model, max_tokens=6000, system=system,
                                             messages=[{"role": "user", "content": user}])
        except Exception as e:  # noqa: BLE001
            if isinstance(system, list) and sc._is_cache_error(e):
                m = self._client.messages.create(model=self.model, max_tokens=6000,
                                                 system=sc._strip_cache_control(system),
                                                 messages=[{"role": "user", "content": user}])
            else:
                raise
        u = m.usage
        self.usages.append({"in": u.input_tokens, "out": u.output_tokens,
                            "cr": getattr(u, "cache_read_input_tokens", 0) or 0,
                            "cw": getattr(u, "cache_creation_input_tokens", 0) or 0})
        return "".join(b.text for b in m.content if b.type == "text")


def select_20():
    raw = list(csv.DictReader(RAW.open(encoding="utf-8-sig")))
    bench = list(csv.DictReader(BENCH.open(encoding="utf-8-sig")))
    raw_by = {}
    for r in raw:
        s = slug(r.get("linkedin url"))
        if s and s not in raw_by:
            raw_by[s] = r

    def sc_(b):
        try:
            return int(b.get("Lead Score") or 0)
        except ValueError:
            return 0

    matched = [b for b in bench if slug(b.get("LinkedIn URL")) in raw_by]
    matched.sort(key=sc_, reverse=True)
    cat = lambda b: (b.get("Priority Status") or "").strip()
    high = [b for b in matched if cat(b) in ("A+ / Hot", "A / High")][:5]
    Bs = [b for b in matched if cat(b) == "B / Normal"]
    med = [Bs[i * (len(Bs) // 5)] for i in range(5)]
    Cs = [b for b in matched if cat(b) == "C / Low"]
    low = [Cs[len(Cs) // 4 + i * (len(Cs) // 12)] for i in range(5)]   # mid-band C
    lowest = Cs[-5:]                                                    # bottom C
    picked, seen, final = high + med + low + lowest, set(), []
    for b in picked:
        s = slug(b["LinkedIn URL"])
        if s not in seen:
            seen.add(s)
            final.append((s, b, raw_by[s]))
    # top up if dedupe dropped any
    for b in matched:
        if len(final) >= 20:
            break
        s = slug(b["LinkedIn URL"])
        if s not in seen:
            seen.add(s)
            final.append((s, b, raw_by[s]))
    return final[:20], len(raw), len(raw[0]), len(bench), len(bench[0])


def main():
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    out = ROOT / "outputs" / "pilots" / f"internal_alpha_{ts}"
    out.mkdir(parents=True, exist_ok=True)
    checks = {}

    selected, raw_rows, raw_cols, bench_rows, bench_cols = select_20()
    assert len(selected) == 20

    print(f"SOURCE: {RAW.name} ({raw_rows}x{raw_cols}) | BENCHMARK (comparison only): {BENCH.name} ({bench_rows}x{bench_cols})")
    print("Selected 20 (5 high / 5 medium / 5 low / 5 lowest):")
    for i, (s, b, _) in enumerate(selected):
        print(f"  [{i:2d}] old={b['Lead Score']:>3} {b['Priority Status']:<12} {b['First Name']} {b['Last Name']}")

    # (1) ICP extraction
    icp = extract_icp(ICP_PDF)
    checks["1_icp_extraction"] = icp.n_pages > 0 and len(icp.text) > 1000
    # (2) ICP profile creation
    profile = sc.build_scoring_profile("FinTech", icp.text)
    checks["2_icp_profile"] = len(profile.scoring_dimensions) >= 1 and len(profile.category_thresholds) >= 1
    print(f"ICP: {icp.name} ({icp.n_pages}p) | dims={len(profile.scoring_dimensions)} | enrichment_guarded={profile.enrichment_required_fields}")

    # (3) Lead normalization
    leads = [sc.normalize_lead(rr, i) for i, (_, _, rr) in enumerate(selected)]
    checks["3_normalization"] = all(l.fields.get("job_title") for l in leads)
    # (4) Evidence extraction
    checks["4_evidence"] = all(len(ev.extract_evidence(l.fields, l.previous_roles)) > 0 for l in leads)

    # capture engine logs (retries / parser failures)
    buf = io.StringIO()
    h = logging.StreamHandler(buf)
    logging.getLogger("qualification").setLevel(logging.INFO)
    logging.getLogger("qualification").addHandler(h)

    client = CapturingClient()
    print(f"Client={type(client).__name__} real | model={client.model} | batch_size=5")
    results = sc.score_leads(leads, icp.text, "FinTech", client=client, batch_size=5)   # single run
    logtext = buf.getvalue()
    by = {r.lead_index: r for r in results}
    by_lead = {l.index: l for l in leads}
    pairs = [(by_lead[r.lead_index], r) for r in results]

    ok = [r for r in results if r.error is None]
    failed = [r for r in results if r.error is not None]
    # (5) proposal parsing / (6) python decision / (7) guard / (8) confidence+coverage
    checks["5_parsing"] = len(ok) == 20
    checks["6_python_decision"] = all(r.provisional_priority and r.raw_icp_score is not None for r in ok)
    enr = set(profile.enrichment_required_fields)
    leak = [r.lead_index for r in ok if any(r.dimensions.get(n) and r.dimensions[n].points > 0 for n in enr)]
    checks["7_enrichment_guard"] = (leak == [])
    checks["8_confidence_coverage"] = all(r.evidence_coverage is not None and r.decision_confidence is not None for r in ok)

    # (9) workbook + (10) CSV fallbacks
    qdf = export.build_qualified_dataframe(pairs)
    adf = export.build_approved_dataframe(qdf)
    wb_bytes = export.to_workbook_bytes(pairs, "FinTech")
    (out / "internal_alpha.xlsx").write_bytes(wb_bytes)
    (out / "qualified_leads.csv").write_bytes(export.to_qualified_csv_bytes(qdf))
    (out / "approved_for_outreach.csv").write_bytes(export.to_approved_csv_bytes(adf))

    wb = load_workbook(io.BytesIO(wb_bytes))
    ws1 = wb["Qualified Leads"]
    header = [c.value for c in ws1[1]]
    url_col = header.index("LinkedIn URL") + 1
    wbv = {
        "exactly_three_sheets": wb.sheetnames == ["Qualified Leads", "Approved for Outreach", "Summary"],
        "qualified_has_20": ws1.max_row == 21,
        "approved_header_only": wb["Approved for Outreach"].max_row == 1,
        "reviewer_fields_empty": bool((qdf["Human Decision"] == "").all() and (qdf["Rejection Reason"] == "").all() and (qdf["Reviewer Comment"] == "").all()),
        "review_status_pending": bool((qdf["Review Status"] == "Pending").all()),
        "mock_real_distinguishable": set(qdf["Mock Result"]) <= {"MOCK/OFFLINE", "real"} and "real" in set(qdf["Mock Result"]),
        "hyperlinks_work": ws1.cell(row=2, column=url_col).hyperlink is not None,
        "confirmed_suspected_separate": "Confirmed Dealbreakers" in header and "Suspected Dealbreakers" in header,
        "unknown_fields_preserved": bool(qdf["Unknown Fields"].astype(str).str.len().sum() > 0),
        "column_order_ok": header == export.QUALIFIED_COLUMNS,
    }
    checks["9_workbook"] = all(wbv.values())
    reread = list(csv.DictReader(io.StringIO((out / "qualified_leads.csv").read_text(encoding="utf-8-sig"))))
    checks["10_csv_fallback"] = len(reread) == 20 and export.build_approved_dataframe(qdf).empty

    # summary metrics match
    summary = export.build_summary_rows(pairs, "FinTech", ts)
    sd = {k: v for k, v in summary if k}
    checks["summary_matches"] = sd.get("Total processed leads") == 20 and sd.get("Successful leads") == len(ok)

    # cost
    tin = sum(u["in"] for u in client.usages); tout = sum(u["out"] for u in client.usages)
    tcr = sum(u["cr"] for u in client.usages); tcw = sum(u["cw"] for u in client.usages)
    cost = tin/1e6*P_IN + tout/1e6*P_OUT + tcr/1e6*P_CR + tcw/1e6*P_CW
    avg = cost/len(ok) if ok else 0
    n_calls = len(client.usages)
    retries = max(0, n_calls - 4)          # 4 batches expected (20/5)
    parser_failures = logtext.count("no JSON array") + logtext.count("JSONDecodeError")

    from collections import Counter
    dist = Counter(r.provisional_priority for r in ok)

    accept = all(checks.values())
    lines = []
    lines.append("Release 0.3 Internal Alpha — End-to-End Acceptance Test")
    lines.append("=" * 56)
    lines.append(f"Run: {ts}")
    lines.append(f"ACCEPTANCE: {'PASS' if accept else 'FAIL'}")
    lines.append("")
    lines.append("Workflow stage checks:")
    for k, v in checks.items():
        lines.append(f"  [{'PASS' if v else 'FAIL'}] {k}")
    lines.append("")
    lines.append("Workbook verification:")
    for k, v in wbv.items():
        lines.append(f"  [{'PASS' if v else 'FAIL'}] {k}")
    lines.append("")
    lines.append(f"Leads: selected=20 successful={len(ok)} failed={len(failed)}")
    lines.append(f"API calls={n_calls} retries={retries} parser_failures={parser_failures}")
    lines.append(f"Tokens: input={tin} output={tout} cache_read={tcr} cache_write={tcw}")
    lines.append(f"Cost: total=${cost:.4f} avg/successful=${avg:.5f} projected {PROJECTION_N:,}=${avg*PROJECTION_N:,.2f}")
    lines.append("")
    lines.append("Result distribution (new priority):")
    for k, v in dist.most_common():
        lines.append(f"  {k}: {v}")
    lines.append("")
    lines.append("Note: human approval is required before any outreach. Linked Helper import is manual.")
    report = "\n".join(lines)
    (out / "summary.txt").write_text(report + "\n", encoding="utf-8")

    (out / "results.json").write_text(json.dumps({
        "selected": [{"i": i, "url": b["LinkedIn URL"], "name": f"{b['First Name']} {b['Last Name']}",
                      "old_score": b["Lead Score"], "old_cat": b["Priority Status"],
                      "new_priority": by[i].provisional_priority, "raw": by[i].raw_icp_score,
                      "coverage": by[i].evidence_coverage, "confidence": by[i].decision_confidence,
                      "dealbreaker_state": by[i].dealbreaker_state,
                      "confirmed": by[i].confirmed_dealbreakers, "suspected": by[i].suspected_dealbreakers,
                      "unknowns": by[i].unknowns, "mock": by[i].is_mock, "error": by[i].error}
                     for i, (s, b, rr) in enumerate(selected)],
        "checks": checks, "workbook_verification": wbv, "usages": client.usages,
        "cost": {"input": tin, "output": tout, "cache_read": tcr, "cache_write": tcw,
                 "total_usd": cost, "avg_per_ok": avg, "projected_20000": avg * PROJECTION_N},
    }, indent=2, ensure_ascii=False), encoding="utf-8")

    print("\n" + report)
    print(f"\nOutputs written to: {out}")
    print("ACCEPTANCE:", "PASS" if accept else "FAIL")


if __name__ == "__main__":
    main()
